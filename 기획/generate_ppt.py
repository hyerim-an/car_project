import os
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1
slide_layout = prs.slide_layouts[0] # Title Slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "노후 설비의 완벽한 부활:\n초경량 딥러닝 예지보전 대시보드 도입 제안"
subtitle.text = "차체 조립 공정 노후 SPOT 용접기 기반 지능형 품질 관제 시스템"

# Slide 2
slide_layout = prs.slide_layouts[1] # Title and Content
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Contents"
content = slide.placeholders[1]
content.text = "1. Problem: 사후 검사의 굴레와 암묵지 유실의 위기\n2. Target: 비용과 생산성 사이에서 고뇌하는 현장과 관리자\n3. Solution: 비용과 복잡함을 덜어낸 지능형 예지보전 파트너\n4. Core Tech: PyTorch 기반 비지도 학습 이상 탐지 알고리즘\n5. Core UX & MVP: 1초 만에 확인하는 직관적 대시보드\n6. Business Value & ROI: 다운타임 20% 단축 및 스마트 팩토리 확장"

# Slide 3
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Problem: 사후 검사의 굴레와 암묵지 유실의 위기"
content = slide.placeholders[1]
content.text = "• 샘플링 검사 한계: 조립 후 불량 발견 시 대규모 재작업 및 글로벌 리콜 리스크 상존\n• 고비용 인프라 장벽: 설비 전면 교체 및 엔터프라이즈 DB 구축을 위한 막대한 예산 확보 불가\n• 기술 전수 단절: 전극팁 마모 및 미세 전류 판별이 베테랑 작업자의 감(암묵지)에 전적으로 의존"

# Slide 4
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Target: 비용과 생산성 사이에서 고뇌하는 현장과 관리자"
content = slide.placeholders[1]
content.text = "• 의사결정자 (생산운영팀장):\n  \"막대한 투자 예산 없이 생산 라인 가동률을 방어하고, 리콜 사태를 철저히 막아내야 한다.\"\n• 실사용자 (현장 엔지니어):\n  \"매일 수많은 컬럼이 혼재된 엑셀/CSV 원시 데이터를 수기로 파싱하고 분석하는 소모적 작업에서 벗어나고 싶다.\""

# Slide 5
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Solution: 지능형 스마트 예지보전 파트너"
content = slide.placeholders[1]
content.text = "• 초경량 인프라 (Lightweight): 고비용 DB를 배제하고 FastAPI와 React+Vite 기반의 가볍고 빠른 웹 아키텍처 적용\n• 자동 맵핑 파이프라인: 수많은 데이터 중 핵심 변수 4종(가압력, 전류, 전압, 통전시간)만 자동 파싱\n• 암묵지의 자산화: 베테랑의 직감을 데이터 기반의 객관적인 탐지 시스템으로 전환하여 체계 구축"

# Slide 6
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Core Tech: PyTorch 기반 비지도 학습 이상 탐지"
content = slide.placeholders[1]
content.text = "• 비지도 학습 모델 (Autoencoder): 불량 라벨 데이터가 부족한 현실을 고려, 정상 데이터의 패턴만 집중 학습하여 기준선(Baseline) 구축\n• 재구성 오류 (Reconstruction Error): 모델에 입력된 데이터와 복원된 데이터 간의 차이(Loss)를 계산하여 임계값 초과 시 즉각 '이상'으로 판별\n• 3대 불량 세부 분류 로직: 단순 알람을 넘어 탐지된 이상치를 [1: 파임불량, 2: 용접부족, 3: 크랙발생] 코드로 구체화하여 제공"

# Slide 7
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Core UX & MVP 시연: 1초 만에 확인하는 직관적 대시보드"
content = slide.placeholders[1]
content.text = "• 산업용 다크 모드 테마: 현장 조명 및 작업자의 눈 피로도를 최소화하며 이상치 표출(네온 레드/오렌지) 가시성 극대화\n• MVP 시연 시나리오 (End-to-End 플로우)\n  1. 원시 파일 업로드: 복잡한 다차원 CSV 파일을 드래그 앤 드롭\n  2. 자동 파싱: 업로드 즉시 백엔드에서 4대 핵심 변수 추출 완료\n  3. 실시간 대시보드 갱신: 불량 유형별 마커 색상 변경 및 Alerts Log에 상세 원인 동시 표출"

# Slide 8
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Business Value & ROI: 다운타임 20% 단축과 미래 확장"
content = slide.placeholders[1]
content.text = "• 비용 효율 최적화: 고가의 기존 엔터프라이즈 솔루션 예상 도입 비용 대비 5% 미만 수준으로 인프라 구축\n• 생산성 극대화 보장: 섀도우 테스트를 통해 설비 다운타임 20% 단축 및 보고서 작성 리드타임 1시간 이내 축소\n• 스마트 팩토리 확장 (Next Step): E-GMP 기반 알루미늄 및 이종 소재 접합 공정 데이터베이스 조기 확보 및 디지털 트윈 연계 기반 마련"

# Slide 9
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Q&A"
content = slide.placeholders[1]
content.text = "질의응답 및 토론"

ppt_path = r"c:\Users\hlahn\car_project\기획\예지보전_시스템_제안서.pptx"
prs.save(ppt_path)
print("PPTX saved to", ppt_path)
